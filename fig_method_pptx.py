"""Editable PowerPoint version of the method figure (fig_method).
Every box/arrow/label is a native shape so it can be moved in PowerPoint.

Slide 1: horizon-gated architecture. Slide 2: gate schedule curve.
Usage: python3 fig_method_pptx.py  ->  fig_method.pptx
"""
import numpy as np
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

# palette (matches figures_paper.py)
BLUE, AQUA, YELLOW, VIOLET, RED = "2a78d6", "1baf7a", "eda100", "4a3aa7", "e34948"
INK, MUTED, GRID = "1f2937", "6b7280", "d5d8dc"
ENCF, TGTF, GATEF = "dbe9f9", "f4f4f2", "fdf1d8"


def rgb(h):
    return RGBColor.from_string(h)


prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(5.6)
blank = prs.slide_layouts[6]


def box(slide, x, y, w, h, text, fill, line=MUTED, size=13, bold=False,
        tcolor=INK, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sp = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    sp.fill.solid(); sp.fill.fore_color.rgb = rgb(fill)
    sp.line.color.rgb = rgb(line); sp.line.width = Pt(1)
    sp.shadow.inherit = False
    tf = sp.text_frame; tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = rgb(tcolor)
    return sp


def label(slide, x, y, w, text, size=10, color=MUTED, bold=False,
          align=PP_ALIGN.CENTER, italic=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(0.4))
    tf = tb.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = align
    for i, ln in enumerate(text.split("\n")):
        para = p if i == 0 else tf.add_paragraph()
        para.alignment = align
        r = para.add_run(); r.text = ln
        r.font.size = Pt(size); r.font.color.rgb = rgb(color)
        r.font.bold = bold; r.font.italic = italic
    return tb


def arrow(slide, x0, y0, x1, y1, color=INK, width=1.5, double=False):
    cn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                                    Inches(x0), Inches(y0), Inches(x1), Inches(y1))
    cn.line.color.rgb = rgb(color); cn.line.width = Pt(width)
    le = cn.line._get_or_add_ln()
    from pptx.oxml.ns import qn
    tail = le.makeelement(qn("a:tailEnd"), {"type": "triangle"})
    le.append(tail)
    if double:
        head = le.makeelement(qn("a:headEnd"), {"type": "triangle"})
        le.append(head)
    return cn


# ---------------- Slide 1: architecture ----------------
s = prs.slides.add_slide(blank)
label(s, 0.3, 0.15, 9, "A   Horizon-gated latent prediction", size=17, bold=True,
      color=INK, align=PP_ALIGN.LEFT)

# row: past -> encoder -> z blocks
box(s, 0.4, 1.15, 1.9, 1.0, "past patches\nx≤t", TGTF)
arrow(s, 2.3, 1.65, 3.0, 1.65)
box(s, 3.0, 1.15, 1.9, 1.0, "causal\nencoder", ENCF)
arrow(s, 4.9, 1.65, 5.7, 1.5)

# z = [slow | fast]
box(s, 5.7, 1.05, 1.0, 0.85, "z_slow", BLUE, line="ffffff", tcolor="ffffff",
     bold=True, size=12, shape=MSO_SHAPE.RECTANGLE)
box(s, 6.7, 1.05, 2.6, 0.85, "z_fast", YELLOW, line="ffffff", tcolor=INK,
     bold=True, size=12, shape=MSO_SHAPE.RECTANGLE)
label(s, 5.7, 0.62, 3.6, "z_t ∈ R⁶⁴   (16 | 48 dims)", size=10,
      color=MUTED, align=PP_ALIGN.LEFT)

# dcor between blocks
arrow(s, 6.7, 1.98, 6.5, 1.98, color=VIOLET, width=1.5, double=True)
label(s, 5.9, 2.05, 1.8, "dcor λ", size=10, color=VIOLET, align=PP_ALIGN.LEFT)

# gate on fast path
box(s, 7.8, 3.0, 1.5, 0.95, "× g(Δ)\ngate", GATEF)
arrow(s, 8.4, 1.9, 8.5, 3.0, color=INK)

# predictor
box(s, 3.4, 3.3, 2.4, 1.0, "predictor\nP(·, Δ)", ENCF)
arrow(s, 6.1, 1.9, 4.9, 3.3, color=BLUE)          # slow path always open
arrow(s, 7.9, 3.5, 5.8, 3.7, color=INK)           # gated fast path
label(s, 5.9, 2.75, 2.0, "slow path:\nalways open", size=9, color=BLUE,
      align=PP_ALIGN.LEFT)

# target + loss
box(s, 0.4, 3.3, 1.9, 1.0, "EMA target\nz̄_{t+Δ}", TGTF)
arrow(s, 3.4, 3.8, 2.3, 3.8, color=INK)
label(s, 2.25, 3.28, 1.2, "L₂ /\nInfoNCE", size=9, color=MUTED)

# takeaway caption
label(s, 0.4, 4.55, 9.2,
      "Long horizons see only z_slow ⇒ slow factors are assigned there "
      "(Prop. 1);\nbottleneck + dcor exclude fast factors (Prop. 2).",
      size=11, color=INK, align=PP_ALIGN.LEFT)

# ---------------- Slide 2: gate schedule ----------------
s2 = prs.slides.add_slide(blank)
label(s2, 0.3, 0.15, 9.4, "B   Gate schedule: fast block fades beyond τ",
      size=17, bold=True, color=INK, align=PP_ALIGN.LEFT)

# plot area in inches
ox, oy, pw, ph = 1.1, 1.2, 7.9, 3.2      # origin x/y (top-left of axes box), width, height
xmax = 140.0
def px(d):  return ox + pw * d / xmax
def py(g):  return oy + ph * (1 - g)     # g in [0,1]

# axes
arrow(s2, ox, oy + ph, ox + pw + 0.15, oy + ph, color=MUTED, width=1.25)  # x
cn = s2.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(ox), Inches(oy),
                             Inches(ox), Inches(oy + ph))
cn.line.color.rgb = rgb(MUTED); cn.line.width = Pt(1.25)

# trained-horizon gridlines
for off in [1, 4, 16, 64, 128]:
    gl = s2.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(px(off)),
                                 Inches(oy), Inches(px(off)), Inches(oy + ph))
    gl.line.color.rgb = rgb(GRID); gl.line.width = Pt(0.75)
    label(s2, px(off) - 0.25, oy - 0.32, 0.5, str(off), size=9, color=MUTED)

# sigmoid gate curve as a freeform of small segments
tau, w = 16.0, 4.0
ds = np.linspace(0, xmax, 120)
gs = 1 / (1 + np.exp((ds - tau) / w))
from pptx.oxml.ns import qn
fb = s2.shapes.build_freeform(Emu(int(Inches(px(ds[0])))), Emu(int(Inches(py(gs[0])))))
fb.add_line_segments([(Emu(int(Inches(px(d)))), Emu(int(Inches(py(g)))))
                       for d, g in zip(ds[1:], gs[1:])], close=False)
curve = fb.convert_to_shape()
curve.fill.background()
curve.line.color.rgb = rgb(YELLOW); curve.line.width = Pt(2.5)

# tau marker
tl = s2.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(px(tau)), Inches(oy),
                             Inches(px(tau)), Inches(oy + ph))
tl.line.color.rgb = rgb(INK); tl.line.width = Pt(1)
tl.line._get_or_add_ln().append(
    tl.line._get_or_add_ln().makeelement(qn("a:prstDash"), {"val": "dash"}))
label(s2, px(tau) + 0.1, oy + 0.55, 2.4, "τ = c·T_ac\n(label-free)",
      size=11, color=INK, align=PP_ALIGN.LEFT)

# axis labels
label(s2, ox, oy + ph + 0.35, pw, "horizon Δ  (patches; trained set marked)",
      size=11, color=INK)
label(s2, ox - 0.95, oy - 0.05, 0.9, "1", size=9, color=MUTED, align=PP_ALIGN.RIGHT)
label(s2, ox - 0.95, oy + ph - 0.18, 0.9, "0", size=9, color=MUTED, align=PP_ALIGN.RIGHT)
vt = s2.shapes.add_textbox(Inches(ox - 1.0), Inches(oy + ph/2 - 0.2), Inches(1.4), Inches(0.4))
vp = vt.text_frame.paragraphs[0]; vp.alignment = PP_ALIGN.CENTER
vr = vp.add_run(); vr.text = "gate g(Δ)"; vr.font.size = Pt(11); vr.font.color.rgb = rgb(INK)
vt.rotation = 270

prs.save("fig_method.pptx")
print("saved fig_method.pptx (2 slides, native editable shapes)")
